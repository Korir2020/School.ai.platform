"""Check-only endpoint for the Migration Centre. Saves nothing."""
import csv
import io
import json
import re

from drf_spectacular.utils import extend_schema, OpenApiResponse, OpenApiTypes
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from schools.models import Enrollment, School, Stream, Student

from .migration_rules import analyse_upload, detect_header_row
from .student_import import _admin
from .views import _scope

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


class UploadError(Exception):
    pass


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _delimiter(content):
    for line in content.splitlines():
        if line.strip():
            counts = {d: line.count(d) for d in ("\t", ";", ",")}
            best = max(counts, key=counts.get)
            return best if counts[best] else ","
    return ","


def _parse_text(content):
    """Split pasted or uploaded text into rows, whatever the separator."""
    lines = [ln for ln in content.splitlines() if ln.strip() and not re.fullmatch(r"[\s|:+\-=_]+", ln)]
    sample = lines[:6]
    best, top = None, 0
    for d in ("\t", "|", ";", ","):
        count = sum(ln.count(d) for ln in sample)
        if count > top:
            best, top = d, count
    if best is None:
        return [[c.strip() for c in re.split(r"\s{2,}", ln.strip())] for ln in lines]
    if best == "|":
        return [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    return [[c.strip() for c in r] for r in csv.reader(lines, delimiter=best)]


def _trim(rows):
    while rows and not any(rows[-1]):
        rows.pop()
    if len(rows) > MAX_ROWS + 15:
        raise UploadError(f"Too many rows. Upload at most {MAX_ROWS} learners at a time.")
    return rows


def _parse_xlsx(upload):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
        rows = []
        for row in wb.active.iter_rows(values_only=True):
            rows.append([_text(c) for c in row])
            if len(rows) > MAX_ROWS + 20:
                break
        wb.close()
    except Exception:
        raise UploadError("Could not read this Excel file. Save it again as .xlsx or .csv.")
    return _trim(rows)


PDF_PAGE_MARK = re.compile(r"^\s*(page\s*)?\d+\s*(of|/)\s*\d+\s*$", re.I)
MAX_PDF_PAGES = 60


def _parse_pdf(upload):
    raw = upload.read()
    if not raw.startswith(b"%PDF"):
        raise UploadError("This does not look like a PDF file.")
    rows, scanned = [], []
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UploadError("This PDF is password protected. Remove the password and try again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UploadError("This PDF has more than %d pages. Split it and upload it in parts." % MAX_PDF_PAGES)
        for n, page in enumerate(reader.pages, start=1):
            text = page.extract_text(extraction_mode="layout") or ""
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            if not lines:
                try:
                    has_picture = len(page.images) > 0
                except Exception:
                    has_picture = True
                if has_picture:
                    scanned.append(n)
                continue
            for ln in lines:
                if not PDF_PAGE_MARK.match(ln):
                    rows.append([c.strip() for c in re.split(r"\t|\s{2,}", ln)])
    except UploadError:
        raise
    except Exception:
        raise UploadError("Could not read this PDF. Save it again, or use an Excel or CSV file.")
    if scanned:
        nums = ", ".join(str(n) for n in scanned)
        label = "Page %s of this PDF is a picture, not typed text," if len(scanned) == 1 else "Pages %s of this PDF are pictures, not typed text,"
        raise UploadError(label % nums + " so it cannot be read yet. Nothing was skipped. Use an Excel or CSV file, or wait for photo reading.")
    if not rows:
        raise UploadError("No typed text was found in this PDF.")
    head = detect_header_row(rows)
    if head is not None:
        hdr = [c.lower() for c in rows[head]]
        rows = rows[:head + 1] + [r for r in rows[head + 1:] if [c.lower() for c in r] != hdr]
        bad = [i for i in range(head + 1, len(rows)) if len(rows[i]) != len(hdr)]
        if bad:
            shown = "; ".join("line %d: %s" % (i + 1, " | ".join(rows[i])[:50]) for i in bad[:3])
            raise UploadError("Some rows have a different number of columns than the header (" + shown + "). This usually means an empty cell, and a PDF cannot show which column it belonged to. Nothing was added. Fill every cell, or use an Excel or CSV file.")
    return _trim(rows)


def parse_table(upload=None, text=""):
    if upload is not None:
        if upload.size > MAX_BYTES:
            raise UploadError("The file must be 5 MB or smaller.")
        name = (getattr(upload, "name", "") or "").lower()
        if name.endswith(".xlsx"):
            return _parse_xlsx(upload)
        if name.endswith(".pdf"):
            return _parse_pdf(upload)
        if name.endswith(".xls"):
            raise UploadError("Old .xls files are not supported. Save as .xlsx or .csv and try again.")
        raw = upload.read()
    else:
        raw = (text or "").encode("utf-8")
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = raw.decode("cp1252", errors="replace")
    if not content.strip():
        raise UploadError("The file is empty.")
    try:
        rows = [[_text(c) for c in r] for r in _parse_text(content)]
    except csv.Error:
        raise UploadError("Could not read this file as a table.")
    return _trim(rows)


def _confirmations(raw):
    if not raw:
        return None
    try:
        data = json.loads(raw) if isinstance(raw, str) else raw
    except ValueError:
        raise UploadError("The confirmations were not valid.")
    if not isinstance(data, dict):
        raise UploadError("The confirmations were not valid.")
    out = {}
    for key, codes in data.items():
        k = int(key) if str(key).isdigit() else ("all" if key == "all" else None)
        if k is None or not isinstance(codes, list):
            raise UploadError("The confirmations were not valid.")
        out[k] = [str(c) for c in codes]
    return out


FIX_FIELDS = ("first_name", "last_name", "admission_number", "assessment_number", "class_level", "stream")


def _fixes(raw):
    if not raw:
        return None
    try:
        data = json.loads(raw) if isinstance(raw, str) else raw
    except ValueError:
        raise UploadError("The corrections were not valid.")
    if not isinstance(data, dict) or len(data) > MAX_ROWS:
        raise UploadError("The corrections were not valid.")
    out = {}
    for key, fields in data.items():
        if not str(key).isdigit() or not isinstance(fields, dict):
            raise UploadError("The corrections were not valid.")
        clean = {}
        for name, value in fields.items():
            if name not in FIX_FIELDS or not isinstance(value, str) or len(value) > 120:
                raise UploadError("The corrections were not valid.")
            clean[name] = value.strip()
        if clean:
            out[int(key)] = clean
    return out or None


def build_structure(school_id):
    levels = {}
    streams = Stream.objects.filter(school_id=school_id, is_active=True, class_level__is_active=True).select_related("class_level")
    for s in streams:
        cl = s.class_level
        entry = levels.setdefault(cl.id, {"id": cl.id, "name": cl.name, "level_number": cl.level_number, "streams": []})
        entry["streams"].append({"id": s.id, "name": s.name})
    enrolled = (Enrollment.objects
                .filter(student__school_id=school_id, is_active=True, class_level__is_active=True)
                .values("class_level_id", "class_level__name", "class_level__level_number").distinct())
    for e in enrolled:
        levels.setdefault(e["class_level_id"], {"id": e["class_level_id"], "name": e["class_level__name"],
                                                "level_number": e["class_level__level_number"], "streams": []})
    return sorted(levels.values(), key=lambda x: (x["level_number"], x["name"]))


def build_existing(school_id):
    rows = Student.objects.filter(school_id=school_id).values("id", "first_name", "last_name", "admission_number", "assessment_number")
    return [{**r, "upi": ""} for r in rows]


@extend_schema(request=OpenApiTypes.OBJECT, responses={
    200: OpenApiResponse(description="Review table. Nothing is saved."),
    400: OpenApiResponse(description="The file could not be read."),
    403: OpenApiResponse(description="Administrator authorization required.")})
@api_view(["POST"])
@parser_classes([MultiPartParser, FormParser])
def migration_check(request):
    school_id, error = _scope(request)
    if error:
        return error
    if not _admin(request.user):
        return Response({"detail": "Only school administrators can check an upload."}, status=403)
    if school_id is None:
        raw = request.data.get("school_id")
        if not str(raw or "").isdigit() or not School.objects.filter(pk=int(raw)).exists():
            return Response({"detail": "Choose a school."}, status=400)
        school_id = int(raw)
    upload = request.FILES.get("file")
    text = request.data.get("text", "")
    if upload is None and not str(text).strip():
        return Response({"detail": "Attach a file or paste the table."}, status=400)
    source = request.data.get("source", "digital")
    if source not in ("digital", "paper"):
        return Response({"detail": "Unknown source."}, status=400)
    try:
        matrix = parse_table(upload, text)
        confirmations = _confirmations(request.data.get("confirmations"))
        fixes = _fixes(request.data.get("fixes"))
    except UploadError as exc:
        return Response({"detail": str(exc)}, status=400)
    result = analyse_upload(
        matrix, build_structure(school_id), existing=build_existing(school_id), source=source,
        name_order=request.data.get("name_order") or None,
        default_class=str(request.data.get("default_class") or ""),
        confirmations=confirmations, fixes=fixes)
    result["committed"] = False
    return Response(result)
